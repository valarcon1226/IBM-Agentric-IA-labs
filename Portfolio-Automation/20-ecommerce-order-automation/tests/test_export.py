"""Pruebas de exportación a archivo Microsoft Excel (.xlsx) de pedidos consolidados."""

import asyncio
import io

import openpyxl

from app.database import get_repository
from app.models import (
    Customer,
    Order,
    OrderItem,
    OrderSource,
    OrderStatus,
    PaymentStatus,
    ShippingAddress,
)


def test_export_orders_to_excel(authenticated_client):
    """Verifica que el endpoint /admin/orders/export genere un archivo .xlsx válido y con datos reales."""
    repo = get_repository()

    order1 = Order(
        id="ORD-EXP-001",
        external_id="101",
        source=OrderSource.WOOCOMMERCE,
        status=OrderStatus.LISTO_PARA_DESPACHAR,
        payment_status=PaymentStatus.PAID,
        customer=Customer(
            name="Mariana Duque",
            email="mariana@example.com",
            phone="+573001112233",
            shipping_address=ShippingAddress(street="Calle 100 # 15-20", city="Bogotá"),
        ),
        items=[
            OrderItem(
                sku="CF-HUI-250",
                title="Café Huila Especial (250g)",
                quantity=2,
                unit_price=28000.0,
                total_price=56000.0,
            )
        ],
        total_amount=56000.0,
    )

    order2 = Order(
        id="ORD-EXP-002",
        external_id="102",
        source=OrderSource.SHOPIFY,
        status=OrderStatus.ENTREGADO,
        payment_status=PaymentStatus.PAID,
        customer=Customer(
            name="Felipe Castro",
            email="felipe@example.com",
            phone="+573105556677",
            shipping_address=ShippingAddress(street="Carrera 43A # 1-50", city="Medellín"),
        ),
        items=[
            OrderItem(
                sku="CF-GEI-250",
                title="Café Geisha Sierra Nevada (250g)",
                quantity=1,
                unit_price=55000.0,
                total_price=55000.0,
            )
        ],
        total_amount=55000.0,
    )

    asyncio.run(repo.create_order(order1))
    asyncio.run(repo.create_order(order2))

    # Descargar Excel
    response = authenticated_client.get("/admin/orders/export")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]

    # Leer archivo con openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Pedidos Café de Origen"

    # Verificar encabezados
    headers = [cell.value for cell in ws[1]]
    assert "ID Pedido" in headers
    assert "Plataforma" in headers
    assert "Cliente" in headers
    assert "Total" in headers

    # Verificar filas
    assert ws.max_row == 3  # Encabezado + 2 filas
    row1_vals = [cell.value for cell in ws[2]]
    assert "ORD-EXP-001" in row1_vals or "ORD-EXP-002" in row1_vals


def test_export_with_filter(authenticated_client):
    """Verifica que los filtros de estado se respeten al exportar a Excel."""
    repo = get_repository()

    order1 = Order(
        id="ORD-FIL-001",
        external_id="1",
        source=OrderSource.WOOCOMMERCE,
        status=OrderStatus.LISTO_PARA_DESPACHAR,
        payment_status=PaymentStatus.PAID,
        customer=Customer(name="A", email="a@example.com", shipping_address=ShippingAddress()),
        items=[],
        total_amount=100.0,
    )
    order2 = Order(
        id="ORD-FIL-002",
        external_id="2",
        source=OrderSource.SHOPIFY,
        status=OrderStatus.ENTREGADO,
        payment_status=PaymentStatus.PAID,
        customer=Customer(name="B", email="b@example.com", shipping_address=ShippingAddress()),
        items=[],
        total_amount=200.0,
    )

    asyncio.run(repo.create_order(order1))
    asyncio.run(repo.create_order(order2))

    # Filtrar solo "entregado"
    response = authenticated_client.get("/admin/orders/export?status=entregado")
    assert response.status_code == 200

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.max_row == 2  # Encabezado + 1 fila filtrada
    row_vals = [cell.value for cell in ws[2]]
    assert "ORD-FIL-002" in row_vals
