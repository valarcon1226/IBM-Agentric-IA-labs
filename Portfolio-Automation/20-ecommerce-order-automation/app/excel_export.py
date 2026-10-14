"""Generador de reportes de pedidos en formato Microsoft Excel (.xlsx) usando openpyxl."""

import io
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import Order


def export_orders_to_excel(orders: List[Order]) -> bytes:
    """Genera un archivo Excel (.xlsx) en memoria con la lista de pedidos normalizados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedidos Café de Origen"

    # Estilos
    header_fill = PatternFill(start_color="3E2723", end_color="3E2723", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    headers = [
        "ID Pedido",
        "Plataforma",
        "ID Externo",
        "Fecha",
        "Estado",
        "Estado Pago",
        "Cliente",
        "Email",
        "Teléfono",
        "Dirección",
        "Ciudad",
        "Departamento",
        "Total",
        "Moneda",
        "Ítems del Pedido",
    ]

    ws.append(headers)

    # Formato de fila encabezado
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Filas de datos
    for row_idx, order in enumerate(orders, start=2):
        items_str = "; ".join([f"{it.quantity}x {it.title}" for it in order.items])
        fecha_str = order.created_at.strftime("%Y-%m-%d %H:%M")

        row_values = [
            order.id,
            order.source.value.capitalize(),
            order.external_id,
            fecha_str,
            order.status.value.replace("_", " ").title(),
            order.payment_status.value.capitalize(),
            order.customer.name,
            order.customer.email,
            order.customer.phone,
            order.customer.shipping_address.street,
            order.customer.shipping_address.city,
            order.customer.shipping_address.state,
            order.total_amount,
            order.currency,
            items_str,
        ]
        ws.append(row_values)

        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 3, 4, 5, 6, 14):
                cell.alignment = center_align
            elif col_idx == 13:
                cell.alignment = right_align
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = left_align

    # Ajustar ancho de columnas automáticamente
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()
