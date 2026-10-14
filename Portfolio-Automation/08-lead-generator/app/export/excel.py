"""Excel export (DECISIONES-08 L5): "Leads" / "Resumen" / "Fuentes" sheets."""

from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import Lead, SourceStat

_LEADS_HEADERS = [
    "Empresa",
    "Categoría",
    "Teléfono",
    "Email",
    "Email válido",
    "Sitio web",
    "Dirección",
    "Ciudad",
    "Puntaje",
    "Fuente",
]

_SCORE_BUCKETS = (("Bajo (0-39)", 0, 39), ("Medio (40-69)", 40, 69), ("Alto (70-100)", 70, 100))


def _style_header(sheet: Worksheet, headers: list[str]) -> None:
    bold = Font(bold=True)
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = bold
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(row=1, column=len(headers)).coordinate}"


def _write_leads_sheet(sheet: Worksheet, leads: list[Lead]) -> None:
    _style_header(sheet, _LEADS_HEADERS)

    for row, lead in enumerate(leads, start=2):
        sheet.cell(row=row, column=1, value=lead.company_name)
        sheet.cell(row=row, column=2, value=lead.category)
        sheet.cell(row=row, column=3, value=lead.phone_e164 or lead.phone_raw)

        email_cell = sheet.cell(row=row, column=4, value=lead.email)
        if lead.email:
            email_cell.hyperlink = f"mailto:{lead.email}"
            email_cell.style = "Hyperlink"

        sheet.cell(row=row, column=5, value=lead.is_email_valid)

        website_cell = sheet.cell(row=row, column=6, value=lead.website)
        if lead.website:
            url = lead.website if lead.website.startswith("http") else f"https://{lead.website}"
            website_cell.hyperlink = url
            website_cell.style = "Hyperlink"

        sheet.cell(row=row, column=7, value=lead.address_raw)
        sheet.cell(row=row, column=8, value=lead.city)
        sheet.cell(row=row, column=9, value=lead.lead_score)
        sheet.cell(row=row, column=10, value=lead.source_name)

    last_row = max(len(leads) + 1, 2)
    score_range = f"I2:I{last_row}"
    sheet.conditional_formatting.add(
        score_range,
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="F8696B",
            mid_type="num",
            mid_value=50,
            mid_color="FFEB84",
            end_type="num",
            end_value=100,
            end_color="63BE7B",
        ),
    )

    widths = {
        "A": 32,
        "B": 18,
        "C": 16,
        "D": 28,
        "E": 12,
        "F": 32,
        "G": 36,
        "H": 14,
        "I": 10,
        "J": 16,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width


def _write_resumen_sheet(sheet: Worksheet, leads: list[Lead]) -> None:
    _style_header(sheet, ["Categoría", "Total"])
    by_category = Counter(lead.category or "Sin categoría" for lead in leads)
    row = 2
    for category, total in sorted(by_category.items()):
        sheet.cell(row=row, column=1, value=category)
        sheet.cell(row=row, column=2, value=total)
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Rango de puntaje").font = Font(bold=True)
    sheet.cell(row=row, column=2, value="Total").font = Font(bold=True)
    row += 1
    for label, low, high in _SCORE_BUCKETS:
        count = sum(1 for lead in leads if low <= lead.lead_score <= high)
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=count)
        row += 1

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 12


def _write_fuentes_sheet(sheet: Worksheet, sources: list[SourceStat]) -> None:
    _style_header(sheet, ["Fuente", "Leads encontrados", "Duplicados fusionados"])
    for row, stat in enumerate(sources, start=2):
        sheet.cell(row=row, column=1, value=stat.source_name)
        sheet.cell(row=row, column=2, value=stat.leads_found)
        sheet.cell(row=row, column=3, value=stat.duplicates_merged)

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 20


def write_leads_excel(leads: list[Lead], sources: list[SourceStat], path: str | Path) -> None:
    workbook = Workbook()
    leads_sheet = workbook.active
    leads_sheet.title = "Leads"
    _write_leads_sheet(leads_sheet, leads)

    _write_resumen_sheet(workbook.create_sheet("Resumen"), leads)
    _write_fuentes_sheet(workbook.create_sheet("Fuentes"), sources)

    workbook.save(path)
