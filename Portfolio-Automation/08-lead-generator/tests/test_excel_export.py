from pathlib import Path

import openpyxl

from app.export.excel import write_leads_excel
from app.schemas import Lead, SourceStat


def _sample_leads() -> list[Lead]:
    return [
        Lead(
            company_name="Restaurante El Buen Sabor",
            category="restaurantes",
            website="elbuensabor.com.co",
            email="contacto@elbuensabor.com.co",
            is_email_valid=True,
            phone_e164="+576012345678",
            address_raw="Cra 7 # 45-10",
            city="Bogotá",
            lead_score=90,
            source_name="demo-directory",
        ),
        Lead(
            company_name="Asadero La Brasa Dorada",
            category="restaurantes",
            phone_e164="+573012345678",
            city="Bogotá",
            lead_score=20,
            source_name="demo-directory",
        ),
        Lead(
            company_name="Contadores Ruiz SAS",
            category="contadores",
            website="contadoresruiz.com.co",
            city="Bogotá",
            lead_score=60,
            source_name="demo-directory",
        ),
    ]


def test_write_leads_excel_produces_the_three_sheets(tmp_path: Path):
    path = tmp_path / "leads.xlsx"
    write_leads_excel(_sample_leads(), [SourceStat("demo-directory", 3, 1)], path)

    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == ["Leads", "Resumen", "Fuentes"]


def test_leads_sheet_has_bold_headers_filter_and_links(tmp_path: Path):
    path = tmp_path / "leads.xlsx"
    write_leads_excel(_sample_leads(), [SourceStat("demo-directory", 3, 1)], path)

    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Leads"]

    header_row = [cell.value for cell in sheet[1]]
    assert header_row == [
        "Empresa",
        "Categoría",
        "Teléfono",
        "Email",
        "Email válido",
        "Sitio web",
        "Dirección",
        "Ciudad",
        "Puntaje",
        "Fuente",
    ]
    assert all(cell.font.bold for cell in sheet[1])
    assert sheet.auto_filter.ref == "A1:J1"

    email_cell = sheet.cell(row=2, column=4)
    assert email_cell.hyperlink is not None
    assert email_cell.hyperlink.target == "mailto:contacto@elbuensabor.com.co"

    website_cell = sheet.cell(row=2, column=6)
    assert website_cell.hyperlink is not None


def test_leads_sheet_has_a_conditional_color_scale_on_the_score_column(tmp_path: Path):
    path = tmp_path / "leads.xlsx"
    write_leads_excel(_sample_leads(), [SourceStat("demo-directory", 3, 1)], path)

    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Leads"]

    rules = sheet.conditional_formatting
    found = False
    for rng in rules:
        if str(rng.sqref) == "I2:I4":
            for rule in rng.rules:
                if rule.type == "colorScale":
                    found = True
    assert found, "expected a colorScale conditional formatting rule on the score column"


def test_resumen_sheet_has_category_and_score_range_totals(tmp_path: Path):
    path = tmp_path / "leads.xlsx"
    write_leads_excel(_sample_leads(), [SourceStat("demo-directory", 3, 1)], path)

    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Resumen"]
    values = [row[0].value for row in sheet.iter_rows(min_row=2)]
    assert "contadores" in values
    assert "restaurantes" in values
    assert "Alto (70-100)" in values


def test_fuentes_sheet_has_the_source_row(tmp_path: Path):
    path = tmp_path / "leads.xlsx"
    write_leads_excel(_sample_leads(), [SourceStat("demo-directory", 3, 1)], path)

    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Fuentes"]
    assert sheet.cell(row=2, column=1).value == "demo-directory"
    assert sheet.cell(row=2, column=2).value == 3
    assert sheet.cell(row=2, column=3).value == 1
