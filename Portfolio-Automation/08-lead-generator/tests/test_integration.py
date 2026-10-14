import asyncio
import io

import httpx
import pytest
from openpyxl import load_workbook
from playwright.async_api import async_playwright

API_BASE = "http://localhost:8008"
DEMO_DIR = "http://localhost:8088"
API_KEY = "change_me_api_key"

HEADERS = {"Authorization": f"Bearer {API_KEY}"}


def is_stack_up():
    try:
        r = httpx.get(f"{API_BASE}/health", timeout=2.0)
        return r.status_code == 200
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not is_stack_up(), reason="Stack is not up")


@pytest.mark.e2e
@pytest.mark.integration
def test_full_job_pipeline():
    # 1. Start job
    job_payload = {
        "keyword": "",
        "location": "BogotÃ¡",
        "sources": ["demo-directory"],
        "max_results": 100,
    }

    response = httpx.post(
        f"{API_BASE}/api/v1/jobs", json=job_payload, headers=HEADERS, timeout=10.0
    )
    assert response.status_code == 201, f"Failed to start job: {response.text}"
    job_data = response.json()
    job_id = job_data["id"]

    # 2. Wait for completion
    import time

    max_retries = 30
    for _ in range(max_retries):
        r = httpx.get(f"{API_BASE}/api/v1/jobs/{job_id}", headers=HEADERS, timeout=10.0)
        assert r.status_code == 200
        status_data = r.json()
        if status_data["status"] == "COMPLETED":
            break
        elif status_data["status"] == "FAILED":
            pytest.fail("Job failed")
        time.sleep(2)
    else:
        pytest.fail("Job timed out")

    assert status_data["total_found"] >= 30

    # Dedup check will be done by inspecting the Fuentes sheet from the export

    # 3. Fetch leads export
    export_r = httpx.get(
        f"{API_BASE}/api/v1/leads/export?format=xlsx", headers=HEADERS, timeout=10.0
    )
    assert export_r.status_code == 200

    # 4. Validate Excel
    file_bytes = io.BytesIO(export_r.content)
    wb = load_workbook(file_bytes)
    assert "Leads" in wb.sheetnames

    ws = wb["Leads"]
    header = [cell.value for cell in ws[1]]
    expected_headers = [
        "Empresa",
        "Categoría",
        "Teléfono",
        "Email",
        "Email válido",
        "Sitio web",
        "Dirección",
        "Ciudad",
        "Puntaje",
        "Fuente",
    ]
    assert header == expected_headers
    assert ws.max_row >= 31

    # Check categories and scores
    categories = set()
    for row in range(2, ws.max_row + 1):
        categories.add(ws.cell(row=row, column=2).value)
        assert ws.cell(row=row, column=9).value >= 0  # Puntaje >= 0

    assert "restaurantes" in categories
    assert "contadores" in categories
    assert "clinicas-dentales" in categories

    assert "Fuentes" in wb.sheetnames
    ws_fuentes = wb["Fuentes"]
    total_merged = sum(
        ws_fuentes.cell(row=r, column=3).value or 0 for r in range(2, ws_fuentes.max_row + 1)
    )
    assert total_merged >= 1


@pytest.mark.e2e
@pytest.mark.integration
def test_playwright_demo_directory():
    # A simple Playwright test against the JS listing
    async def run_test():
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            page = await browser.new_page()
            await page.goto(f"{DEMO_DIR}/")

            # Check if page is loaded
            title = await page.title()

            # We assume there's some content we can wait for, such as a list or items
            # Just ensure the page does not throw and we can evaluate JS
            assert "Demo" in title or title != ""

            await browser.close()

    asyncio.run(run_test())
