"""Script para regenerar entregable-plantilla.xlsx con las hojas de 08 + Seguimiento.

Uso:
    python service/generar_plantilla.py
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

_SEGUIMIENTO_HEADERS = [
    "Empresa",
    "Categoría",
    "Teléfono",
    "Email",
    "Estado",
    "Fecha de contacto",
    "Respuesta",
    "Próximo paso / Notas",
]


def _build_base_workbook_if_needed(sample_path: Path, output_path: Path) -> None:
    """Si no existe samples/leads-ejemplo.xlsx, lo genera desde demo_data con el código del 08."""
    if sample_path.exists():
        return

    from app.export.excel import write_leads_excel
    from app.normalization import normalize_name, normalize_phone
    from app.schemas import Lead, SourceStat
    from app.scoring import score_lead
    from tests.fixtures.demo_data import BUSINESSES

    leads: list[Lead] = []
    for b in BUSINESSES:
        norm_phone = normalize_phone(b.phone)
        norm_name = normalize_name(b.name)
        lead = Lead(
            company_name=b.name,
            normalized_name=norm_name,
            category=b.category,
            website=f"demo-directory/sitios/{b.site}/index.html" if b.site else None,
            email=b.site_email,
            is_email_valid=bool(b.site_email and "@" in b.site_email),
            phone_raw=b.phone,
            phone_e164=norm_phone,
            address_raw=b.address,
            city="Bogotá",
            lat=4.6097,
            lng=-74.0817,
            rating=4.5,
            source_url=f"demo-directory/{b.category}/index.html",
            lead_score=0,
            source_name="demo-directory",
        )
        lead.lead_score = score_lead(lead)
        leads.append(lead)

    sources = [
        SourceStat(
            source_name="demo-directory",
            leads_found=len(leads),
            duplicates_merged=1,
        )
    ]
    write_leads_excel(leads, sources, sample_path)


def generate_template(output_path: Path | None = None) -> Path:
    sample_path = BASE_DIR / "samples" / "leads-ejemplo.xlsx"
    if output_path is None:
        output_path = Path(__file__).resolve().parent / "entregable-plantilla.xlsx"

    _build_base_workbook_if_needed(sample_path, output_path)

    wb = openpyxl.load_workbook(sample_path)

    # Si ya existe la hoja Seguimiento (de una corrida previa), la eliminamos para reconstruirla limpia
    if "Seguimiento" in wb.sheetnames:
        del wb["Seguimiento"]

    ws_leads = wb["Leads"]
    ws_seg = wb.create_sheet("Seguimiento")

    # Encabezados
    bold_font = Font(bold=True, color="1F2937")
    header_fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")

    for col_idx, header in enumerate(_SEGUIMIENTO_HEADERS, start=1):
        cell = ws_seg.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws_seg.freeze_panes = "A2"
    ws_seg.auto_filter.ref = f"A1:{get_column_letter(len(_SEGUIMIENTO_HEADERS))}1"

    # Población con los leads existentes
    sample_statuses = [
        (
            "Respondió - Interesado",
            "2026-10-10",
            "Solicitó propuesta y llamada de demo",
            "Reunión agendada para el viernes 10:00 AM",
        ),
        (
            "Contactado (Email 1)",
            "2026-10-12",
            "Abierto (sin respuesta aún)",
            "Enviar secuencia de seguimiento #2 en 3 días",
        ),
        (
            "Contactado (Email 2)",
            "2026-10-08",
            "No por el momento (presupuesto Q1)",
            "Volver a contactar en febrero 2027",
        ),
    ]

    max_r = ws_leads.max_row
    for r in range(2, max_r + 1):
        empresa = ws_leads.cell(row=r, column=1).value
        categoria = ws_leads.cell(row=r, column=2).value
        telefono = ws_leads.cell(row=r, column=3).value
        email = ws_leads.cell(row=r, column=4).value

        idx = r - 2
        if idx < len(sample_statuses):
            estado, fecha, resp, nota = sample_statuses[idx]
        else:
            estado = "1. Pendiente de contacto"
            fecha = None
            resp = "Sin contactar"
            nota = None

        row_dest = r
        ws_seg.cell(row=row_dest, column=1, value=empresa)
        ws_seg.cell(row=row_dest, column=2, value=categoria)
        ws_seg.cell(row=row_dest, column=3, value=telefono)

        cell_email = ws_seg.cell(row=row_dest, column=4, value=email)
        if email:
            cell_email.hyperlink = f"mailto:{email}"
            cell_email.style = "Hyperlink"

        ws_seg.cell(row=row_dest, column=5, value=estado)
        ws_seg.cell(row=row_dest, column=6, value=fecha)
        ws_seg.cell(row=row_dest, column=7, value=resp)
        ws_seg.cell(row=row_dest, column=8, value=nota)

    # Anchos de columna
    col_widths = {
        "A": 34,  # Empresa
        "B": 20,  # Categoría
        "C": 18,  # Teléfono
        "D": 30,  # Email
        "E": 26,  # Estado
        "F": 18,  # Fecha de contacto
        "G": 32,  # Respuesta
        "H": 40,  # Próximo paso / Notas
    }
    for col_letter, width in col_widths.items():
        ws_seg.column_dimensions[col_letter].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> None:
    out = generate_template()
    wb = openpyxl.load_workbook(out)
    print(f"Plantilla generada exitosamente en: {out}")
    print(f"Hojas: {wb.sheetnames}")
    for name in wb.sheetnames:
        print(f"  - {name}: {wb[name].max_row} filas")


if __name__ == "__main__":
    main()
